"""真实 Edge 页面验证导入模板、预检查、错误下载及确认后记录读回。"""

from __future__ import annotations

import io
import json
import secrets
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

PROJECT = Path(__file__).resolve().parents[1]
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
URL = "http://127.0.0.1:15173/"


def excel(template, rows):
    book = openpyxl.load_workbook(io.BytesIO(template))
    for line, values in enumerate(rows, 3):
        for column, value in enumerate(values, 1):
            book.active.cell(line, column, value)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def main():
    suffix = secrets.token_hex(4)
    screenshot = PROJECT / "ai" / "evidence" / "r15-07-import-browser.png"
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=str(EDGE), headless=True)
        context = browser.new_context(viewport={"width": 1540, "height": 1050}, locale="zh-CN", accept_downloads=True)
        page = context.new_page()
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto(URL)
        page.get_by_role("button", name="还没有账号？立即注册").click()
        page.get_by_placeholder("例如 admin01").fill("ui_import_" + suffix)
        page.get_by_placeholder("你的姓名或团队名称").fill("导入验收")
        page.get_by_placeholder("至少 8 位").fill("UiImportCheck9!" + suffix)
        page.get_by_role("button", name="注册并进入").click()
        page.get_by_role("button", name="创建第一个系统").click()
        page.get_by_placeholder("例如：采购管理").fill("导入浏览器系统")
        page.get_by_placeholder("例如：purchase").fill("ui_import_" + suffix)
        page.get_by_role("button", name="创建并进入").click()
        page.get_by_role("heading", name="模块组", exact=True).wait_for()
        page.get_by_title("创建模块组").click()
        page.get_by_placeholder("例如：销售业务").fill("销售业务")
        page.get_by_placeholder("例如：sales").fill("sales")
        page.get_by_role("button", name="创建模块组", exact=True).click()
        page.get_by_title("创建业务模块").click()
        page.get_by_placeholder("例如：客户").fill("客户")
        page.get_by_placeholder("例如：customer").fill("customer")
        page.get_by_role("button", name="创建模块", exact=True).click()
        page.get_by_role("button", name="添加字段").click()
        page.get_by_placeholder("例如：客户名称").fill("客户名称")
        page.get_by_placeholder("例如：customer_name").fill("customer_name")
        page.get_by_role("button", name="保存字段").click()
        page.get_by_role("button", name="发布模块", exact=True).click()
        page.get_by_text("发布成功，当前为第 1 版。").wait_for()
        page.get_by_role("button", name="运行预览").click()
        page.locator(".record-toolbar").wait_for()
        page.wait_for_load_state("networkidle")
        page.get_by_role("button", name="导入 Excel").click()
        panel = page.locator(".export-panel").filter(has_text="导入 Excel")
        with page.expect_download() as info:
            panel.get_by_role("link", name="下载模板").click()
        template = Path(info.value.path()).read_bytes()
        assert openpyxl.load_workbook(io.BytesIO(template)).active.cell(2, 2).value == "customer_name"
        panel.locator("input[type=file]").set_input_files({
            "name": "invalid.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": excel(template, [(None, "可用客户"), ("R_BAD", None)])})
        panel.get_by_text("预检查：1 / 2 行通过").wait_for()
        with page.expect_download() as info:
            panel.get_by_role("link", name="下载错误行").first.click()
        error_book = openpyxl.load_workbook(io.BytesIO(Path(info.value.path()).read_bytes()))
        assert error_book["错误原因"].cell(2, 2).value
        assert error_book.active.cell(2, 2).value == "customer_name"
        assert page.locator(".record-row").count() == 0
        panel.locator("input[type=file]").set_input_files({
            "name": "valid.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": excel(template, [(None, "客户一"), (None, "客户二")])})
        panel.get_by_text("预检查：2 / 2 行通过").wait_for()
        panel.get_by_role("button", name="确认整批导入").click()
        page.get_by_text("已导入 2 行。").wait_for()
        page.locator(".record-row", has_text="客户一").wait_for()
        page.locator(".record-row", has_text="客户二").wait_for()
        panel.get_by_text("已导入 2 行").wait_for()
        page.screenshot(path=str(screenshot), full_page=True)
        assert not errors, errors
        browser.close()
    print(json.dumps({"status": "passed", "screenshot": str(screenshot),
                      "checks": ["template_download", "invalid_preview", "error_download",
                                 "ready_preview", "confirm", "record_readback", "no_page_errors"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
