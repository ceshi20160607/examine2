"""用隔离系统的 Excel 模板验证模块导入预演、执行和错误行。需要 openpyxl。"""

import csv
import hashlib
import importlib.util
import io
import json
import secrets
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location("fixture", Path(__file__).with_name("verify-external-flow.py"))
fixture = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixture)


def raw(path, token, method="GET", content=None, headers=None):
    values = {"Authorization": "Bearer " + token, **(headers or {})}
    request = urllib.request.Request(fixture.BASE + path, data=content, headers=values, method=method)
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            return response.status, response.read()
    except urllib.error.HTTPError as error:
        return error.code, error.read()


def expect_error(path, method, body, token, status, code):
    actual, response = fixture.request(path, method, body, token)
    if actual != status or response.get("code") != code:
        raise AssertionError(f"{method} {path}: expected {status}/{code}, got {actual}/{response}")


def upload_xlsx(content, token, account_id):
    sha = hashlib.sha256(content).hexdigest()
    session = fixture.expected("/api/files/uploads", "POST", {
        "originalName": "g6-customers.xlsx",
        "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "expectedSize": len(content), "expectedSha256": sha}, token)
    status, chunk = raw(f"/api/files/uploads/{session['id']}/chunks?offset=0", token, "PUT", content, {
        "Content-Type": "application/octet-stream", "X-Upload-Token": session["uploadToken"],
        "X-Chunk-Sha256": sha})
    if status != 200 or json.loads(chunk).get("code") != "OK":
        raise AssertionError(f"Excel 分片上传失败：{status}/{chunk[:300]}")
    status, completed = raw(f"/api/files/uploads/{session['id']}/complete", token, "POST", b"", {
        "X-Upload-Token": session["uploadToken"]})
    result = json.loads(completed)
    if status != 200 or result.get("code") != "OK":
        raise AssertionError(f"Excel 完成上传失败：{status}/{result}")
    file = result["data"]
    if file["scanStatus"] != "CLEAN":
        raise AssertionError(f"Excel 文件未通过安全扫描：{file}")
    fixture.expected(f"/api/files/{file['id']}/references", "POST", {
        "ownerType": "ACCOUNT", "ownerId": str(account_id),
        "fieldCode": "managed_files", "referenceType": "DOCUMENT"}, token)
    return file["id"]


def main():
    suffix = secrets.token_hex(4)

    def register(name):
        username = f"g6_import_{name}_{suffix}"
        result = fixture.expected("/api/auth/register", "POST", {
            "username": username, "password": "correct-password", "displayName": f"G6 导入{name}",
            "email": username + "@example.com", "systemName": f"G6 导入{name}系统",
            "systemCode": f"g6_import_{name}_{suffix}"})
        return username, result

    owner_name, owner = register("owner")
    member_name, member = register("member")
    reader_name, _ = register("reader")
    token = owner["tokens"]["accessToken"]
    system_id, tenant_id = owner["systemId"], owner["tenantId"]
    group = fixture.expected("/api/admin/module-config/groups", "POST", {
        "code": "g6_customer_group", "name": "客户管理", "sortOrder": 10}, token)
    module = fixture.expected("/api/admin/module-config/modules", "POST", {
        "groupId": group["id"], "code": "g6_import_customer", "name": "客户"}, token)
    module_id = module["module"]["id"]
    fixture.expected(f"/api/admin/module-config/modules/{module_id}/fields", "POST", {
        "code": "customer_code", "name": "客户编码", "fieldType": "TEXT", "required": False,
        "uniqueValue": True, "searchable": True, "sortOrder": 10, "config": {}}, token)
    draft = fixture.expected(f"/api/admin/module-config/modules/{module_id}/draft", token=token)
    fixture.expected(f"/api/admin/module-config/modules/{module_id}/publish", "POST", {
        "expectedDraftRevision": draft["module"]["draftRevision"]}, token)

    def role(name, importer):
        actions = ["LIST", "DETAIL"] + (["CREATE", "UPDATE", "IMPORT"] if importer else [])
        permissions = [{"resourceType": "MODULE", "resourceCode": "g6_import_customer",
                        "actionCode": action,
                        "dataScopeType": "ALL" if action in ("CREATE", "IMPORT") else "SELF",
                        "dataScopeJson": "{}"}
                       for action in actions]
        if importer:
            permissions += [{"resourceType": "FILE", "resourceCode": "OBJECT", "actionCode": action,
                             "dataScopeType": "ALL", "dataScopeJson": "{}"}
                            for action in ("UPLOAD", "VIEW", "REFERENCE", "DOWNLOAD")]
        policies = [{"resourceCode": "g6_import_customer", "fieldCode": "customer_code",
                     "channel": channel, "readable": True, "writable": importer,
                     "maskStrategy": None} for channel in ("PAGE", "FILE")]
        saved = fixture.expected("/api/admin/system/authorization/roles", "POST", {
            "code": name, "name": name, "permissions": permissions,
            "fieldPolicies": policies}, token)
        return fixture.expected(f"/api/admin/system/authorization/roles/{saved['id']}/publish", "POST", {
            "reason": "G6 模块导入权限验证", "expectedVersion": saved["version"]}, token)

    import_role = role("g6_importer", True)
    read_role = role("g6_import_reader", False)
    importer = fixture.expected("/api/admin/system/authorization/members", "POST", {
        "account": member_name, "roleIds": [import_role["id"]]}, token)
    reader = fixture.expected("/api/admin/system/authorization/members", "POST", {
        "account": reader_name, "roleIds": [read_role["id"]]}, token)
    original = fixture.expected("/api/runtime/modules/g6_import_customer/records", "POST", {
        "title": "导入前客户", "status": "ACTIVE", "ownerMemberId": importer["systemMemberId"],
        "participantMemberIds": [], "fields": {"customer_code": "EX-001"}}, token)
    other = fixture.expected("/api/runtime/modules/g6_import_customer/records", "POST", {
        "title": "其他负责人客户", "status": "ACTIVE", "ownerMemberId": reader["systemMemberId"],
        "participantMemberIds": [], "fields": {"customer_code": "OTHER-001"}}, token)

    def enter(username):
        platform = fixture.expected("/api/auth/login", "POST", {
            "username": username, "password": "correct-password"})
        system = fixture.expected(f"/api/systems/{system_id}/enter", "POST", {}, platform["accessToken"])
        return system["tokens"]["accessToken"]

    member_token = enter(member_name)
    member_account_id = member["accountId"]
    reader_token = enter(reader_name)
    base = "/api/runtime/modules/g6_import_customer/imports"
    expect_error(base + "/template", "GET", None, reader_token, 403, "IMPORT_PERMISSION_DENIED")
    template = fixture.expected(base + "/template", token=member_token)
    codes = [column["code"] for column in template["columns"]]
    if not {"title", "recordNumber", "status", "customer_code"}.issubset(codes):
        raise AssertionError(f"导入模板缺少目标字段：{template}")
    status, template_bytes = raw(base + "/template.xlsx", member_token)
    if status != 200:
        raise AssertionError("Excel 模板下载失败")
    workbook = load_workbook(io.BytesIO(template_bytes))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    sheet.delete_rows(2, sheet.max_row)
    headers.append("createdByMemberId")
    sheet.cell(1, len(headers), "createdByMemberId")
    records = [
        {"标题": "更新客户", "客户编码": "EX-001", "createdByMemberId": "999999"},
        {"标题": "新增客户", "客户编码": "NEW-002", "createdByMemberId": "999999"},
        {"标题": "", "客户编码": "BAD-003", "createdByMemberId": "999999"},
        {"标题": "越权客户", "客户编码": "OTHER-001", "createdByMemberId": "999999"},
    ]
    for item in records:
        sheet.append([item.get(header, "") for header in headers])
    output = io.BytesIO()
    workbook.save(output)
    source_id = upload_xlsx(output.getvalue(), member_token, member_account_id)
    source = fixture.expected(base + f"/sources/{source_id}/headers", token=member_token)
    if source["format"] != "XLSX" or "createdByMemberId" not in source["headers"]:
        raise AssertionError(f"Excel 源表头读取错误：{source}")
    expect_error(base + "/preview", "POST", {
        "sourceFileId": source_id, "columnMapping": {"标题": "title", "createdByMemberId": "createdByMemberId"},
        "conflictPolicy": "ERROR", "matchFieldCode": "customer_code"},
        member_token, 422, "IMPORT_MAPPING_TARGET_INVALID")
    preview = fixture.expected(base + "/preview", "POST", {
        "sourceFileId": source_id, "columnMapping": {"标题": "title", "客户编码": "customer_code"},
        "conflictPolicy": "UPDATE", "matchFieldCode": "customer_code",
        "mappingName": "客户编码新增更新"}, member_token)
    if preview["status"] != "PREVIEWED" or preview["summary"] != {
            "create": 1, "update": 1, "skip": 0, "error": 2}:
        raise AssertionError(f"混合行预检查结果错误：{preview}")
    errors = {row["rowNumber"]: row["errorCode"] for row in preview["rows"] if row["status"] == "ERROR"}
    if errors != {4: "IMPORT_TITLE_REQUIRED", 5: "UNIQUE_RULE_VIOLATION"}:
        raise AssertionError(f"预检查未拒绝空标题或他人唯一值：{errors}")
    before = fixture.expected("/api/runtime/modules/g6_import_customer/records", token=token)
    if len(before["records"]) != 2:
        raise AssertionError("预检查意外写入业务记录")
    error_status, error_bytes = raw(base + f"/{preview['id']}/errors.csv", member_token)
    error_rows = list(csv.reader(io.StringIO(error_bytes.decode("utf-8-sig"))))
    if error_status != 200 or {row[0] for row in error_rows[1:]} != {"4", "5"}:
        raise AssertionError(f"错误行下载不完整：{error_status}/{error_rows}")

    queued = fixture.expected(base + f"/{preview['id']}/execute", "POST", {
        "version": preview["version"]}, member_token)
    for _ in range(60):
        result = fixture.expected(base + f"/{queued['id']}", token=member_token)
        if result["status"] in ("COMPLETED", "COMPLETED_WITH_ERRORS", "FAILED"):
            break
        time.sleep(0.3)
    if result["status"] != "COMPLETED_WITH_ERRORS" or result["successRows"] != 2 or result["failedRows"] != 2:
        raise AssertionError(f"导入执行或行级结果错误：{result}")
    after = fixture.expected("/api/runtime/modules/g6_import_customer/records", token=token)["records"]
    if len(after) != 3:
        raise AssertionError(f"导入失败行污染了业务数据：{after}")
    updated = fixture.expected(f"/api/runtime/modules/g6_import_customer/records/{original['id']}", token=token)
    unchanged = fixture.expected(f"/api/runtime/modules/g6_import_customer/records/{other['id']}", token=token)
    created = next((row for row in after if row["title"] == "新增客户"), None)
    if updated["title"] != "更新客户" or updated["fields"]["customer_code"] != "EX-001":
        raise AssertionError(f"唯一字段匹配更新失败：{updated}")
    if unchanged["title"] != "其他负责人客户" or not created:
        raise AssertionError("无权记录被更新或新增行缺失")
    created = fixture.expected(f"/api/runtime/modules/g6_import_customer/records/{created['id']}", token=token)
    if created["fields"]["customer_code"] != "NEW-002" or created["ownerMemberId"] == importer["systemMemberId"]:
        raise AssertionError(f"新增行自定义字段或默认负责人错误：{created}")
    if created["createdByMemberId"] != created["ownerMemberId"]:
        raise AssertionError("导入数据的创建人没有按系统所属人补齐")
    history = fixture.expected(base, token=member_token)
    if not any(item["id"] == result["id"] for item in history):
        raise AssertionError("导入历史缺少执行批次")
    expect_error(base + f"/{result['id']}", "GET", None, reader_token, 403, "IMPORT_PERMISSION_DENIED")
    export_base = "/api/runtime/modules/g6_import_customer/exports"
    config = fixture.expected(export_base + "/config", token=token)
    if not any(field["code"] == "customer_code" for field in config["availableFields"]):
        raise AssertionError("导入的自定义字段未进入导出配置")
    saved = fixture.expected(export_base + "/config", "PUT", {
        "name": "导入结果核对", "selectedFields": ["title", "customer_code"],
        "relationDisplay": "LABEL", "expectedVersion": None}, token)
    fixture.expected(export_base + "/config/publish", "POST", {
        "expectedDraftRevision": saved["draftRevision"]}, token)
    export_query = {"lifecycleState": "ACTIVE", "tenantScope": "ALL", "search": "",
                    "filters": [], "sortField": "title", "sortDirection": "ASC",
                    "selectedRecordIds": [], "format": "CSV",
                    "idempotencyKey": "g6-import-export-" + suffix, "selectedFields": []}
    exported = fixture.expected(export_base, "POST", export_query, token)
    for _ in range(60):
        exported = fixture.expected(export_base + f"/{exported['id']}", token=token)
        if exported["status"] in ("COMPLETED", "FAILED"):
            break
        time.sleep(0.3)
    if exported["status"] != "COMPLETED" or exported["exportedRows"] != 3:
        raise AssertionError(f"同模块导入后导出失败：{exported}")
    export_status, export_bytes = raw(export_base + f"/{exported['id']}/download", token)
    export_rows = list(csv.reader(io.StringIO(export_bytes.decode("utf-8-sig"))))
    if export_status != 200 or export_rows[0] != ["标题", "客户编码"] or {
            tuple(row) for row in export_rows[1:]} != {
            ("更新客户", "EX-001"), ("新增客户", "NEW-002"),
            ("其他负责人客户", "OTHER-001")}:
        raise AssertionError(f"导入与导出结果不一致：{export_status}/{export_rows}")
    if "--restart" in sys.argv:
        subprocess.run(["docker", "restart", "unexamine-g1-preview"], check=True, capture_output=True)
        for _ in range(70):
            try:
                with urllib.request.urlopen(fixture.BASE + "/actuator/health", timeout=2) as response:
                    if json.loads(response.read()).get("status") == "UP":
                        break
            except OSError:
                pass
            time.sleep(0.5)
        else:
            raise AssertionError("导入后预览服务重启未恢复")
        persistent = fixture.expected(base + f"/{result['id']}", token=member_token)
        if persistent["status"] != "COMPLETED_WITH_ERRORS" or persistent["successRows"] != 2:
            raise AssertionError("重启后导入批次结果不一致")
        created_after_restart = fixture.expected(
            f"/api/runtime/modules/g6_import_customer/records/{created['id']}", token=token)
        if created_after_restart["fields"]["customer_code"] != "NEW-002":
            raise AssertionError("重启后导入主记录和自定义字段不一致")
        export_restart_status, export_after_restart = raw(
            export_base + f"/{exported['id']}/download", token)
        if export_restart_status != 200 or export_after_restart != export_bytes:
            raise AssertionError("重启后同模块导出文件与原文件不同")
    print(json.dumps({"systemId": system_id, "tenantId": tenant_id,
                      "ownerUsername": owner_name, "importerUsername": member_name,
                      "moduleId": module_id, "sourceFileId": source_id,
                      "batchId": result["id"], "jobId": result["jobId"],
                      "updatedRecordId": original["id"], "createdRecordId": created["id"],
                      "otherOwnerRecordId": other["id"], "preview": preview["summary"],
                      "finalStatus": result["status"], "successRows": result["successRows"],
                      "failedRows": result["failedRows"], "errorRows": errors,
                      "integratedExportBatchId": exported["id"],
                      "integratedExportJobId": exported["jobId"],
                      "integratedExportRows": export_rows,
                      "restartVerified": "--restart" in sys.argv}, ensure_ascii=False))


if __name__ == "__main__":
    main()
