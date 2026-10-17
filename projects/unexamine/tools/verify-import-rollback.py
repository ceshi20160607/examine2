"""隔离系统验证导入文件重复、严格更新和批次安全回滚。需要 openpyxl。"""

import importlib.util
import io
import json
import secrets
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


spec = importlib.util.spec_from_file_location("import_fixture", Path(__file__).with_name("verify-module-import.py"))
import_fixture = importlib.util.module_from_spec(spec)
spec.loader.exec_module(import_fixture)
fixture = import_fixture.fixture


def main():
    suffix = secrets.token_hex(4)
    username = "g6_rollback_owner_" + suffix
    owner = fixture.expected("/api/auth/register", "POST", {
        "username": username, "password": "correct-password", "displayName": "G6 回滚管理员",
        "email": username + "@example.com", "systemName": "G6 回滚验证系统",
        "systemCode": "g6_rollback_" + suffix})
    token = owner["tokens"]["accessToken"]
    system_id, account_id = owner["systemId"], owner["accountId"]
    group = fixture.expected("/api/admin/module-config/groups", "POST", {
        "code": "g6_rollback_group", "name": "客户管理", "sortOrder": 10}, token)
    module = fixture.expected("/api/admin/module-config/modules", "POST", {
        "groupId": group["id"], "code": "g6_rollback_customer", "name": "客户"}, token)
    module_id = module["module"]["id"]
    fixture.expected(f"/api/admin/module-config/modules/{module_id}/fields", "POST", {
        "code": "customer_code", "name": "客户编码", "fieldType": "TEXT", "required": False,
        "uniqueValue": True, "searchable": True, "sortOrder": 10, "config": {}}, token)
    draft = fixture.expected(f"/api/admin/module-config/modules/{module_id}/draft", token=token)
    fixture.expected(f"/api/admin/module-config/modules/{module_id}/publish", "POST", {
        "expectedDraftRevision": draft["module"]["draftRevision"]}, token)
    records_base = "/api/runtime/modules/g6_rollback_customer/records"
    imports_base = "/api/runtime/modules/g6_rollback_customer/imports"
    original = fixture.expected(records_base, "POST", {
        "title": "原始客户", "status": "ACTIVE", "participantMemberIds": [],
        "fields": {"customer_code": "EX-001"}}, token)
    status, template_bytes = import_fixture.raw(imports_base + "/template.xlsx", token)
    if status != 200:
        raise AssertionError("无法取得发布模块 Excel 模板")

    def workbook_source(items, actor_token=token, actor_account_id=account_id):
        workbook = load_workbook(io.BytesIO(template_bytes))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        sheet.delete_rows(2, sheet.max_row)
        for item in items:
            sheet.append([item.get(header, "") for header in headers])
        output = io.BytesIO()
        workbook.save(output)
        return import_fixture.upload_xlsx(output.getvalue(), actor_token, actor_account_id)

    def preview(items, policy, actor_token=token, actor_account_id=account_id):
        source_id = workbook_source(items, actor_token, actor_account_id)
        return fixture.expected(imports_base + "/preview", "POST", {
            "sourceFileId": source_id,
            "columnMapping": {"标题": "title", "客户编码": "customer_code"},
            "conflictPolicy": policy, "matchFieldCode": "customer_code",
            "mappingName": "G6 安全回滚验证"}, actor_token)

    def execute(batch, actor_token=token):
        queued = fixture.expected(imports_base + f"/{batch['id']}/execute", "POST", {
            "version": batch["version"]}, actor_token)
        for _ in range(60):
            result = fixture.expected(imports_base + f"/{queued['id']}", token=actor_token)
            if result["status"] in ("COMPLETED", "COMPLETED_WITH_ERRORS", "FAILED"):
                return result
            time.sleep(0.3)
        raise AssertionError("导入任务未在限时内结束")

    def rollback(batch_id, actor_token=token):
        return fixture.expected(imports_base + f"/{batch_id}/rollback", "POST", {
            "reason": "隔离测试批次安全回滚"}, actor_token)

    strict = preview([
        {"标题": "第一次更新", "客户编码": "EX-001"},
        {"标题": "文件内重复", "客户编码": "EX-001"},
        {"标题": "缺失匹配", "客户编码": "MISSING-001"},
    ], "UPDATE_ONLY")
    errors = {row["rowNumber"]: row["errorCode"] for row in strict["rows"] if row["status"] == "ERROR"}
    if strict["summary"] != {"create": 0, "update": 1, "skip": 0, "error": 2} or errors != {
            3: "IMPORT_DUPLICATE_MATCH_IN_FILE", 4: "IMPORT_MATCH_NOT_FOUND"}:
        raise AssertionError(f"严格更新未识别文件内重复或缺失匹配：{strict['summary']}/{errors}")
    before = fixture.expected(records_base + f"/{original['id']}", token=token)
    if before["title"] != "原始客户":
        raise AssertionError("严格更新预演写入了记录")
    strict_done = execute(strict)
    if strict_done["status"] != "COMPLETED_WITH_ERRORS" or strict_done["successRows"] != 1:
        raise AssertionError(f"严格更新执行错误：{strict_done}")
    full = rollback(strict_done["id"])
    restored = fixture.expected(records_base + f"/{original['id']}", token=token)
    if full["status"] != "ROLLED_BACK" or restored["title"] != "原始客户":
        raise AssertionError("严格更新批次未完整回滚")

    mixed = preview([
        {"标题": "批次更新客户", "客户编码": "EX-001"},
        {"标题": "批次新增客户", "客户编码": "NEW-002"},
    ], "UPDATE")
    if mixed["summary"] != {"create": 1, "update": 1, "skip": 0, "error": 0}:
        raise AssertionError(f"新增并更新预演错误：{mixed['summary']}")
    mixed_done = execute(mixed)
    if mixed_done["status"] != "COMPLETED" or mixed_done["successRows"] != 2:
        raise AssertionError("新增并更新任务未完成")
    created_id = next(row["targetRecordId"] for row in mixed_done["rows"]
                      if row["operation"] == "CREATE")
    after_import = fixture.expected(records_base + f"/{original['id']}", token=token)
    changed = fixture.expected(records_base + f"/{original['id']}", "PUT", {
        "title": "导入后人工修改", "recordNumber": after_import["recordNumber"],
        "status": after_import["status"], "ownerMemberId": after_import["ownerMemberId"],
        "departmentId": after_import["departmentId"],
        "participantMemberIds": after_import["participantMemberIds"],
        "fields": {"customer_code": "EX-001"}, "version": after_import["version"]}, token)
    if changed["title"] != "导入后人工修改":
        raise AssertionError("未能模拟导入后的合法修改")
    partial = rollback(mixed_done["id"])
    row_states = {row["rowNumber"]: (row["status"], row["errorCode"]) for row in partial["rows"]}
    if partial["status"] != "ROLLBACK_PARTIAL" or row_states != {
            2: ("ROLLBACK_CONFLICT", "IMPORT_ROLLBACK_VERSION_CONFLICT"),
            3: ("ROLLED_BACK", None)}:
        raise AssertionError(f"批次部分回滚未保护后续修改：{partial['status']}/{row_states}")
    preserved = fixture.expected(records_base + f"/{original['id']}", token=token)
    created = fixture.expected(records_base + f"/{created_id}", token=token)
    if preserved["title"] != "导入后人工修改" or not created["deleted"]:
        raise AssertionError("部分回滚覆盖了后续修改或没有撤销本批新增")
    repeated = rollback(mixed_done["id"])
    if repeated["status"] != "ROLLBACK_PARTIAL" or repeated["summary"]["rollbackConflicts"] != 1:
        raise AssertionError("再次尝试时丢失未解决的回滚冲突")
    member_name = "g6_rollback_member_" + suffix
    member_account = fixture.expected("/api/auth/register", "POST", {
        "username": member_name, "password": "correct-password", "displayName": "G6 导入成员",
        "email": member_name + "@example.com", "systemName": "G6 成员平台账号",
        "systemCode": "g6_rollback_member_" + suffix})
    policies = [{"resourceCode": "g6_rollback_customer", "fieldCode": "customer_code",
                 "channel": channel, "readable": True, "writable": True, "maskStrategy": None}
                for channel in ("PAGE", "FILE")]

    def permission(action, scope):
        return {"resourceType": "MODULE", "resourceCode": "g6_rollback_customer",
                "actionCode": action, "dataScopeType": scope, "dataScopeJson": "{}"}

    role_permissions = [permission(action, "ALL" if action in ("CREATE", "IMPORT") else "SELF")
                        for action in ("LIST", "DETAIL", "CREATE", "UPDATE", "IMPORT")]
    role_permissions += [{"resourceType": "FILE", "resourceCode": "OBJECT",
                          "actionCode": action, "dataScopeType": "ALL", "dataScopeJson": "{}"}
                         for action in ("UPLOAD", "VIEW", "REFERENCE", "DOWNLOAD")]
    role = fixture.expected("/api/admin/system/authorization/roles", "POST", {
        "code": "g6_rollback_importer", "name": "G6 导入成员",
        "permissions": role_permissions, "fieldPolicies": policies}, token)
    role = fixture.expected(f"/api/admin/system/authorization/roles/{role['id']}/publish", "POST", {
        "reason": "G6 导入回滚权限验证", "expectedVersion": role["version"]}, token)
    fixture.expected("/api/admin/system/authorization/members", "POST", {
        "account": member_name, "roleIds": [role["id"]]}, token)

    def member_token():
        platform = fixture.expected("/api/auth/login", "POST", {
            "username": member_name, "password": "correct-password"})
        return fixture.expected(f"/api/systems/{system_id}/enter", "POST", {},
                                platform["accessToken"])["tokens"]["accessToken"]

    limited_token = member_token()
    member_batch = preview([{"标题": "成员新增客户", "客户编码": "MEM-003"}],
                           "UPDATE", limited_token, member_account["accountId"])
    member_done = execute(member_batch, limited_token)
    if member_done["status"] != "COMPLETED":
        raise AssertionError("成员导入没有完成")
    member_created_id = member_done["rows"][0]["targetRecordId"]
    denied_rollback = rollback(member_done["id"], limited_token)
    if denied_rollback["status"] != "ROLLBACK_PARTIAL" or denied_rollback["rows"][0]["errorCode"] != "PERMISSION_DENIED":
        raise AssertionError("无删除权限的回滚未保留明确拒绝")
    revised = fixture.expected("/api/admin/system/authorization/roles", "POST", {
        "id": role["id"], "code": "g6_rollback_importer", "name": "G6 导入成员",
        "permissions": role_permissions + [permission("DELETE", "ALL")],
        "fieldPolicies": policies, "expectedVersion": role["version"]}, token)
    fixture.expected(f"/api/admin/system/authorization/roles/{revised['id']}/publish", "POST", {
        "reason": "补齐批次回滚删除权限", "expectedVersion": revised["version"]}, token)
    recovered = rollback(member_done["id"], member_token())
    member_created = fixture.expected(records_base + f"/{member_created_id}", token=token)
    if recovered["status"] != "ROLLED_BACK" or not member_created["deleted"] \
            or recovered["rows"][0]["errorCode"] is not None:
        raise AssertionError("修正权限后重试没有撤销本批新增记录")
    if "--restart" in sys.argv:
        subprocess.run(["docker", "restart", "unexamine-g1-preview"], check=True, capture_output=True)
        for _ in range(70):
            try:
                with urllib.request.urlopen(fixture.BASE + "/actuator/health", timeout=2) as response:
                    if json.loads(response.read()).get("status") == "UP":
                        break
            except OSError:
                pass
            time.sleep(0.5)
        else:
            raise AssertionError("重启后预览服务未恢复")
        persisted = fixture.expected(imports_base + f"/{mixed_done['id']}", token=token)
        if persisted["status"] != "ROLLBACK_PARTIAL" or persisted["summary"]["rollbackConflicts"] != 1:
            raise AssertionError("重启后批次回滚结果不一致")
    print(json.dumps({"systemId": system_id, "ownerUsername": username,
                      "moduleId": module_id, "originalRecordId": original["id"],
                      "createdRecordId": created_id, "strictBatchId": strict_done["id"],
                      "strictErrors": errors, "fullRollbackStatus": full["status"],
                      "mixedBatchId": mixed_done["id"], "partialRollbackStatus": partial["status"],
                      "partialRowStates": row_states, "repeatConflictCount": repeated["summary"]["rollbackConflicts"],
                      "memberBatchId": member_done["id"], "permissionDeniedStatus": denied_rollback["status"],
                      "permissionRecoveredStatus": recovered["status"],
                      "restartVerified": "--restart" in sys.argv}, ensure_ascii=False))


if __name__ == "__main__":
    main()
