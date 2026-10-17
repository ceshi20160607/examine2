"""在独立新系统验证 Excel 导入、权限、事务回滚与 MySQL 落库。"""

from __future__ import annotations

import io
import json
import os
import secrets
import urllib.error
import urllib.request
from http.cookiejar import CookieJar

import openpyxl
import pymysql

BASE = "http://127.0.0.1:15173/api"


def browser():
    return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(CookieJar()))


def call(client, method, path, wanted, body=None, content_type="application/json"):
    payload = json.dumps(body).encode() if body is not None and content_type == "application/json" else body
    request = urllib.request.Request(BASE + path, data=payload, method=method,
                                     headers={"Content-Type": content_type})
    try:
        with client.open(request, timeout=20) as response:
            status, raw, headers = response.status, response.read(), response.headers
    except urllib.error.HTTPError as error:
        status, raw, headers = error.code, error.read(), error.headers
    if status != wanted:
        raise AssertionError(f"{method} {path}: wanted {wanted}, got {status}: {raw[:500]}")
    if headers.get_content_type() == "application/json" and raw:
        return json.loads(raw)
    return raw


def workbook(template, rows):
    book = openpyxl.load_workbook(io.BytesIO(template))
    sheet = book.active
    for offset, values in enumerate(rows, 3):
        for column, value in enumerate(values, 1):
            sheet.cell(offset, column, value)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def preview(client, path, mode, excel, wanted=201):
    boundary = "Boundary" + secrets.token_hex(8)
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"mode\"\r\n\r\n{mode}\r\n"
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"import.xlsx\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode()
    body += excel + f"\r\n--{boundary}--\r\n".encode()
    return call(client, "POST", path + "/preview", wanted, body, f"multipart/form-data; boundary={boundary}")


def main():
    suffix = secrets.token_hex(4)
    owner, member, outsider = [browser() for _ in range(3)]
    password = "ImportCheck9!" + suffix
    usernames = ["import_owner_" + suffix, "import_member_" + suffix, "import_other_" + suffix]
    for client, username in zip((owner, member, outsider), usernames):
        call(client, "POST", "/auth/register", 201,
             {"username": username, "password": password, "displayName": username})
    system = call(owner, "POST", "/systems", 201,
                  {"code": "import_" + suffix, "name": "导入验收系统"})
    sid = system["id"]
    base = f"/systems/{sid}"
    role = call(owner, "GET", base + "/roles", 200)[0]
    group = call(owner, "POST", base + "/module-groups", 201, {"code": "business", "name": "业务"})
    module = call(owner, "POST", base + "/modules", 201,
                  {"groupId": group["id"], "code": "orders", "name": "订单"})
    mid = module["id"]
    path = base + f"/modules/{mid}"
    call(owner, "POST", path + "/fields", 201,
         {"key": "title", "label": "订单名称", "type": "TEXT", "required": True, "searchable": True})
    call(owner, "POST", path + "/fields", 201,
         {"key": "amount", "label": "金额", "type": "NUMBER", "required": False, "searchable": True})
    call(owner, "POST", path + "/publish", 200)
    call(owner, "POST", base + f"/roles/{role['id']}/grants", 200,
         {"moduleId": mid, "canView": True, "canCreate": True, "canUpdate": True,
          "canDelete": False, "canExport": False})
    call(owner, "POST", base + "/members/invite", 201,
         {"username": usernames[1], "roleId": role["id"]})
    record_path = path + "/records"
    import_path = path + "/import-jobs"
    template = call(member, "GET", import_path + "/template", 200)
    book = openpyxl.load_workbook(io.BytesIO(template))
    assert [book.active.cell(2, i).value for i in range(1, 4)] == ["record_no", "title", "amount"]
    assert book.active.row_dimensions[2].hidden
    call(outsider, "GET", import_path + "/template", 404)

    invalid = preview(member, import_path, "INSERT", workbook(template,
                      [(None, "有效订单", 12), (None, None, 7)]))
    assert invalid["status"] == "INVALID" and invalid["validRows"] == 1 and invalid["rows"][1]["message"]
    assert call(member, "GET", record_path, 200)["total"] == 0
    errors = call(member, "GET", import_path + f"/{invalid['id']}/errors", 200)
    error_book = openpyxl.load_workbook(io.BytesIO(errors))
    assert error_book["错误原因"].cell(2, 2).value
    assert error_book.active.cell(2, 2).value == "title" and error_book.active.row_dimensions[2].hidden
    assert error_book.active.cell(3, 3).value == "7"
    error_book.active.cell(3, 2, "修正后的订单")
    corrected = io.BytesIO()
    error_book.save(corrected)
    assert preview(member, import_path, "INSERT", corrected.getvalue())["status"] == "READY"
    call(member, "POST", import_path + f"/{invalid['id']}/commit", 409)
    call(outsider, "GET", import_path + f"/{invalid['id']}/errors", 404)

    inserted = preview(member, import_path, "INSERT", workbook(template,
                       [(None, "订单一", 12.5), (None, "订单二", 8)]))
    assert inserted["status"] == "READY" and inserted["validRows"] == 2
    assert call(member, "GET", record_path, 200)["total"] == 0
    result = call(member, "POST", import_path + f"/{inserted['id']}/commit", 200)
    assert result["status"] == "IMPORTED" and result["importedRows"] == 2
    call(member, "POST", import_path + f"/{inserted['id']}/commit", 409)
    own = call(member, "GET", record_path, 200)
    assert own["total"] == 2
    first = next(row for row in own["items"] if row["title"] == "订单一")
    second = next(row for row in own["items"] if row["title"] == "订单二")
    other = call(owner, "POST", record_path, 201, {"values": {"title": "管理员订单"}})

    forbidden = preview(member, import_path, "UPSERT", workbook(template,
                        [(other["recordNo"], "越权修改", 1)]))
    assert forbidden["status"] == "INVALID" and forbidden["rows"][0]["message"]
    valid = preview(member, import_path, "UPSERT", workbook(template,
                    [(first["recordNo"], "订单一更新", 18), (None, "订单三", 3)]))
    assert valid["status"] == "READY" and valid["rows"][0]["action"] == "UPDATE"
    committed = call(member, "POST", import_path + f"/{valid['id']}/commit", 200)
    assert committed["importedRows"] == 2
    assert call(member, "GET", record_path + f"/{first['id']}", 200)["values"]["amount"] == 18
    assert call(member, "GET", record_path, 200)["total"] == 3
    assert call(owner, "GET", record_path + f"/{other['id']}", 200)["title"] == "管理员订单"

    stale = preview(member, import_path, "UPSERT", workbook(template,
                    [(first["recordNo"], "不应写入", 99), (second["recordNo"], "也不应写入", 77)]))
    current = call(member, "GET", record_path + f"/{second['id']}", 200)
    call(member, "PUT", record_path + f"/{second['id']}", 200,
         {"expectedVersion": current["version"], "values": {"title": "中途编辑", "amount": 8}})
    call(member, "POST", import_path + f"/{stale['id']}/commit", 409)
    assert call(member, "GET", record_path + f"/{first['id']}", 200)["title"] == "订单一更新"
    history = call(member, "GET", import_path, 200)
    assert next(job for job in history if job["id"] == stale["id"])["status"] == "FAILED"
    assert all(job["id"] != stale["id"] for job in call(owner, "GET", import_path, 200))

    db = pymysql.connect(host="127.0.0.1", port=3306, user="root",
                         password=os.environ["APP_DB_PASSWORD"], database="unexamine_r15",
                         charset="utf8mb4", cursorclass=pymysql.cursors.DictCursor)
    try:
        with db.cursor() as cursor:
            cursor.execute("SELECT status,imported_rows FROM un_exchange_import_job WHERE id=%s", (valid["id"],))
            assert cursor.fetchone() == {"status": "IMPORTED", "imported_rows": 2}
            cursor.execute("SELECT COUNT(*) AS c FROM un_module_record WHERE module_id=%s", (mid,))
            assert cursor.fetchone()["c"] == 4
            cursor.execute("SELECT version,success FROM flyway_schema_history ORDER BY installed_rank DESC LIMIT 1")
            assert cursor.fetchone() == {"version": "7", "success": 1}
    finally:
        db.close()
    print(json.dumps({"status": "passed", "system_id": sid, "module_id": mid,
                      "checks": ["template", "precheck", "error_xlsx", "no_write_before_commit", "insert",
                                 "idempotent_commit", "upsert", "owner_scope", "atomic_rollback",
                                 "corrected_error_reupload", "history_scope", "mysql_readback", "flyway_v7"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
