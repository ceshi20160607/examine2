"""一个候选系统贯穿建模、成员数据、导入导出、审批消息和权限撤销。"""

from __future__ import annotations

import io
import json
import os
import secrets
from pathlib import Path

import openpyxl
import pymysql
from playwright.sync_api import sync_playwright
from playwright.sync_api import expect

PROJECT = Path(__file__).resolve().parents[1]
EVIDENCE = PROJECT / "ai" / "evidence"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
URL = "http://127.0.0.1:15173/"
PNG = bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000b49444154789c636000020000050001a5f645400000000049454e44ae426082")


def api(context, method, path, wanted, body=None):
    response = context.request.fetch(URL + "api" + path, method=method, data=body)
    if response.status != wanted:
        raise AssertionError(f"{method} {path}: expected {wanted}, got {response.status}: {response.text()[:500]}")
    return response.json() if response.body() and response.headers.get("content-type", "").startswith("application/json") else response.body()


def workbook(template):
    book = openpyxl.load_workbook(io.BytesIO(template))
    assert [book.active.cell(2, column).value for column in (1, 2, 3)] == ["record_no", "title", "amount"]
    book.active.cell(3, 2, "导入采购申请")
    book.active.cell(3, 3, 42.50)
    data = io.BytesIO()
    book.save(data)
    return data.getvalue()


def main():
    suffix = secrets.token_hex(4)
    names = ["journey_owner_" + suffix, "journey_user_" + suffix, "journey_reviewer_" + suffix]
    password = "JourneyCheck9!" + suffix
    checks = []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=EDGE, headless=True)
        contexts = [browser.new_context(viewport={"width": 1540, "height": 1050}, locale="zh-CN", accept_downloads=True) for _ in names]
        owner, applicant, reviewer = contexts
        for context, username in zip(contexts, names):
            api(context, "POST", "/auth/register", 201,
                {"username": username, "password": password, "displayName": username})
        system = api(owner, "POST", "/systems", 201,
                     {"code": "journey_" + suffix, "name": "完整旅程候选系统"})
        sid = system["id"]
        base = f"/systems/{sid}"
        role = api(owner, "GET", base + "/roles", 200)[0]
        group = api(owner, "POST", base + "/module-groups", 201,
                    {"code": "purchase", "name": "采购业务"})
        module = api(owner, "POST", base + "/modules", 201,
                     {"groupId": group["id"], "code": "requests", "name": "采购申请"})
        mid = module["id"]
        path = base + f"/modules/{mid}"
        pages = [context.new_page() for context in contexts]
        owner_page, applicant_page, reviewer_page = pages
        errors = []
        for name, page in zip(names, pages):
            page.on("pageerror", lambda error, name=name: errors.append(f"{name}: {error}"))

        # 管理员在真实页面配置并发布，后续操作只访问这一个模块的发布版本。
        owner_page.goto(URL + f"systems/{sid}")
        owner_page.locator(".module-column .list-item", has_text="采购申请").click()
        for key, label, kind, required, searchable in (
            ("title", "申请标题", "TEXT", True, True),
            ("amount", "申请金额", "MONEY", False, True),
        ):
            owner_page.get_by_role("button", name="添加字段").click()
            owner_page.get_by_placeholder("例如：客户名称").fill(label)
            owner_page.get_by_placeholder("例如：customer_name").fill(key)
            owner_page.get_by_label("字段类型").select_option(kind)
            if required:
                owner_page.locator(".check-row input").nth(0).check()
            if searchable:
                owner_page.locator(".check-row input").nth(1).check()
            owner_page.get_by_role("button", name="保存字段").click()
            owner_page.locator(".modal-backdrop").wait_for(state="detached")
        owner_page.get_by_role("button", name="发布模块", exact=True).click()
        owner_page.get_by_text("发布成功，当前为第 1 版。").wait_for()
        api(owner, "POST", base + f"/roles/{role['id']}/grants", 200,
            {"moduleId": mid, "canView": True, "canCreate": True, "canUpdate": True,
             "canDelete": False, "canExport": True})
        for username in names[1:]:
            api(owner, "POST", base + "/members/invite", 201,
                {"username": username, "roleId": role["id"]})
        members = {member["username"]: member["id"] for member in api(owner, "GET", base + "/members", 200)}
        api(owner, "PUT", path + "/approval", 200,
            {"name": "采购审批", "nodeName": "负责人审核", "approverMemberId": members[names[2]]})
        api(owner, "POST", path + "/approval/publish", 200)
        checks.append("owner_publish_grant_and_approval")

        applicant_page.goto(URL + f"systems/{sid}")
        applicant_page.locator(".runtime-side button", has_text="采购申请").click()
        applicant_page.get_by_role("button", name="新增记录").click()
        applicant_page.locator(".record-form-grid label", has_text="申请标题").locator("input").fill("界面采购申请")
        applicant_page.locator(".record-form-grid label", has_text="申请金额").locator("input").fill("68.50")
        applicant_page.get_by_role("button", name="保存记录").click()
        applicant_page.get_by_text("记录已保存。").wait_for()
        own = api(applicant, "GET", path + "/records", 200)
        own_record = next(item for item in own["items"] if item["title"] == "界面采购申请")
        rid = own_record["id"]
        applicant_page.locator(".record-row", has_text="界面采购申请").click()
        applicant_page.locator(".record-detail").get_by_role("heading", name="界面采购申请").wait_for()
        applicant_page.locator(".record-files input[type=file]").set_input_files(
            {"name": "凭证.png", "mimeType": "image/png", "buffer": PNG})
        applicant_page.locator(".record-files .file-row", has_text="凭证.png").wait_for()
        files = api(applicant, "GET", path + f"/records/{rid}/files", 200)
        assert len(files) == 1
        checks.append("member_create_detail_and_file")

        applicant_page.get_by_role("button", name="导入 Excel").click()
        import_panel = applicant_page.locator(".export-panel").filter(has_text="导入 Excel")
        with applicant_page.expect_download() as download_info:
            import_panel.get_by_role("link", name="下载模板").click()
        template = Path(download_info.value.path()).read_bytes()
        import_panel.locator("input[type=file]").set_input_files(
            {"name": "采购申请.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
             "buffer": workbook(template)})
        import_panel.get_by_text("预检查：1 / 1 行通过").wait_for()
        import_panel.get_by_role("button", name="确认整批导入").click()
        applicant_page.get_by_text("已导入 1 行。").wait_for()
        applicant_page.locator(".record-row", has_text="导入采购申请").wait_for()
        imported = next(item for item in api(applicant, "GET", path + "/records", 200)["items"]
                        if item["title"] == "导入采购申请")
        assert api(applicant, "GET", path + f"/records/{imported['id']}", 200)["values"]["amount"] == 42.5
        checks.append("same_module_excel_import")

        # 第二名成员有相同角色，数据仍按创建人隔离。
        other = api(reviewer, "POST", path + "/records", 201,
                    {"values": {"title": "审阅人采购申请", "amount": 99}})
        applicant_page.locator(".record-search select").select_option("amount")
        applicant_page.get_by_placeholder("字段精确值").fill("42.5")
        applicant_page.get_by_role("button", name="查询").click()
        expect(applicant_page.locator(".record-row")).to_have_count(1)
        applicant_page.locator(".record-row", has_text="导入采购申请").wait_for()
        applicant_page.get_by_role("button", name="导出数据").click()
        export_panel = applicant_page.locator(".export-panel").filter(has_text="导出当前筛选")
        export_panel.get_by_role("button", name="提交导出任务").click()
        applicant_page.get_by_text("导出任务已提交，可在下方刷新状态并下载。").wait_for()
        for _ in range(30):
            export_panel.get_by_role("button", name="刷新状态").click()
            if export_panel.get_by_role("link", name="下载 CSV").count():
                break
            applicant_page.wait_for_timeout(100)
        with applicant_page.expect_download() as download_info:
            export_panel.get_by_role("link", name="下载 CSV").click()
        csv = Path(download_info.value.path()).read_bytes().decode("utf-8-sig")
        assert "导入采购申请" in csv and "审阅人采购申请" not in csv and "界面采购申请" not in csv
        checks.append("filtered_owner_scoped_csv")

        applicant_page.get_by_placeholder("字段精确值").fill("")
        applicant_page.locator(".record-search select").select_option("")
        applicant_page.get_by_role("button", name="查询").click()
        applicant_page.locator(".record-row", has_text="界面采购申请").click()
        applicant_page.get_by_role("button", name="发起审批").click()
        applicant_page.get_by_text("审批已发起，记录处于待审批状态。").wait_for()
        submitted = api(reviewer, "GET", base + "/approval/tasks", 200)
        assert len(submitted) == 1 and submitted[0]["recordId"] == rid
        task_id = submitted[0]["taskId"]
        assert len(api(reviewer, "GET", base + f"/approval/tasks/{task_id}/files", 200)) == 1
        reviewer_page.goto(URL)
        reviewer_page.locator(".system-card", has_text="完整旅程候选系统").click()
        reviewer_page.get_by_role("button", name="消息").click()
        reviewer_page.locator(".message-item", has_text="新的审批待办").click()
        reviewer_page.locator(".access-card", has_text="TASK DETAIL").get_by_role("heading", name="界面采购申请").wait_for()
        reviewer_page.screenshot(path=str(EVIDENCE / "r15-11-integrated-reviewer.png"), full_page=True)
        reviewer_page.get_by_role("button", name="同意", exact=True).click()
        reviewer_page.get_by_text("已同意审批。").wait_for()
        applicant_page.get_by_role("button", name="消息").click()
        applicant_page.locator(".message-item", has_text="审批申请已通过").click()
        applicant_page.locator(".record-detail").get_by_role("heading", name="界面采购申请").wait_for()
        applicant_page.locator(".record-detail", has_text="已通过").wait_for()
        assert api(applicant, "GET", path + f"/records/{rid}", 200)["recordStatus"] == "APPROVED"
        applicant_page.screenshot(path=str(EVIDENCE / "r15-11-integrated-applicant.png"), full_page=True)
        checks.append("approval_message_target_and_attachment")

        # 撤销同一角色权限后，旧消息仍能看到，目标访问和导出必须由接口重新鉴权。
        api(owner, "POST", base + f"/roles/{role['id']}/grants", 200,
            {"moduleId": mid, "canView": False, "canCreate": False, "canUpdate": False,
             "canDelete": False, "canExport": False})
        api(applicant, "GET", path + f"/records/{rid}", 404)
        api(applicant, "GET", path + f"/records/{rid}/files", 404)
        api(applicant, "POST", path + "/export-jobs", 404, {"columnKeys": ["title"]})
        assert len(api(applicant, "GET", base + "/messages", 200)) == 1
        applicant_page.reload()
        applicant_page.get_by_role("button", name="消息").click()
        applicant_page.locator(".message-item", has_text="审批申请已通过").click()
        applicant_page.get_by_text("当前无权打开这条消息对应的模块").wait_for()
        assert not errors, errors
        checks.append("revoked_permission_blocks_target_file_export")

        db = pymysql.connect(host="127.0.0.1", port=3306, user="root",
                             password=os.environ["APP_DB_PASSWORD"], database="unexamine_r15",
                             charset="utf8mb4", cursorclass=pymysql.cursors.DictCursor)
        try:
            with db.cursor() as cursor:
                # 独立连接读取同一系统的记录、导入、导出、审批、消息和文件引用。
                cursor.execute("SELECT id,record_status FROM un_module_record WHERE module_id=%s ORDER BY id", (mid,))
                record_rows = cursor.fetchall()
                assert len(record_rows) == 3 and next(row for row in record_rows if row["id"] == rid)["record_status"] == "APPROVED"
                cursor.execute("SELECT COUNT(*) AS c FROM un_exchange_import_job WHERE module_id=%s AND status='IMPORTED'", (mid,))
                assert cursor.fetchone()["c"] == 1
                cursor.execute("SELECT COUNT(*) AS c FROM un_exchange_job WHERE module_id=%s AND status='SUCCEEDED' AND row_count=1", (mid,))
                assert cursor.fetchone()["c"] == 1
                cursor.execute("SELECT COUNT(*) AS c FROM un_flow_example WHERE record_id=%s AND status='APPROVED'", (rid,))
                assert cursor.fetchone()["c"] == 1
                cursor.execute("SELECT COUNT(*) AS c FROM un_message WHERE system_id=%s", (sid,))
                assert cursor.fetchone()["c"] == 2
                cursor.execute("SELECT COUNT(*) AS c FROM un_file_reference WHERE record_id=%s", (rid,))
                assert cursor.fetchone()["c"] == 1
        finally:
            db.close()
        checks.append("mysql_same_candidate_readback")
        browser.close()
    result = {"status": "passed", "system_id": sid, "module_id": mid,
              "record_ids": [rid, imported["id"], other["id"]], "checks": checks,
              "screenshots": ["r15-11-integrated-reviewer.png", "r15-11-integrated-applicant.png"]}
    (EVIDENCE / "r15-11-integrated-readback.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
