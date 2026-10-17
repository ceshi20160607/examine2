"""在新系统验证常用字段类型从发布、表单写入到筛选、导入和导出。"""

from __future__ import annotations

import io
import json
import os
import secrets
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from http.cookiejar import CookieJar

import openpyxl
import pymysql

BASE = "http://127.0.0.1:15173/api"


def browser():
    return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(CookieJar()))


def call(client, method, path, wanted, body=None, content_type="application/json"):
    payload = json.dumps(body).encode() if body is not None and content_type == "application/json" else body
    request = urllib.request.Request(BASE + path, data=payload, method=method,
                                     headers={"Content-Type": content_type})
    try:
        with client.open(request, timeout=20) as response:
            status, raw, headers = response.status, response.read(), response.headers
    except urllib.error.HTTPError as error:
        status, raw, headers = error.code, error.read(), error.headers
    if status != wanted:
        raise AssertionError(f"{method} {path}: wanted {wanted}, got {status}: {raw[:400]}")
    return json.loads(raw) if raw and headers.get_content_type() == "application/json" else raw


def preview(client, path, excel):
    boundary = "Boundary" + secrets.token_hex(8)
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"mode\"\r\n\r\nINSERT\r\n"
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"fields.xlsx\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode()
    body += excel + f"\r\n--{boundary}--\r\n".encode()
    return call(client, "POST", path + "/preview", 201, body, f"multipart/form-data; boundary={boundary}")


def main():
    suffix = secrets.token_hex(4)
    owner, member = browser(), browser()
    password = "FieldsCheck9!" + suffix
    names = ["fields_owner_" + suffix, "fields_member_" + suffix]
    for client, username in zip((owner, member), names):
        call(client, "POST", "/auth/register", 201,
             {"username": username, "password": password, "displayName": username})
    system = call(owner, "POST", "/systems", 201,
                  {"code": "fields_" + suffix, "name": "字段类型验收系统"})
    sid = system["id"]
    base = f"/systems/{sid}"
    role = call(owner, "GET", base + "/roles", 200)[0]
    group = call(owner, "POST", base + "/module-groups", 201,
                 {"code": "business", "name": "业务"})
    module = call(owner, "POST", base + "/modules", 201,
                  {"groupId": group["id"], "code": "orders", "name": "订单"})
    mid = module["id"]
    path = base + f"/modules/{mid}"
    definitions = [
        ("title", "订单名称", "TEXT", True, True),
        ("description", "说明", "TEXTAREA", False, False),
        ("phone", "手机号", "PHONE", False, True),
        ("email", "邮箱", "EMAIL", False, True),
        ("url", "网址", "URL", False, False),
        ("amount", "金额", "MONEY", False, True),
        ("percent", "百分比", "PERCENT", False, True),
        ("occurred", "发生时间", "DATETIME", False, True),
    ]
    call(owner, "POST", path + "/fields", 400,
         {"key": "bad_description", "label": "错误说明", "type": "TEXTAREA", "required": False, "searchable": True})
    for key, label, kind, required, searchable in definitions:
        call(owner, "POST", path + "/fields", 201,
             {"key": key, "label": label, "type": kind, "required": required, "searchable": searchable})
    call(owner, "POST", path + "/publish", 200)
    call(owner, "POST", base + f"/roles/{role['id']}/grants", 200,
         {"moduleId": mid, "canView": True, "canCreate": True, "canUpdate": True,
          "canDelete": False, "canExport": True})
    call(owner, "POST", base + "/members/invite", 201,
         {"username": names[1], "roleId": role["id"]})
    record_path = path + "/records"
    values = {"title": "客户订单", "description": "第一行\n第二行", "phone": "+8613800138000",
              "email": "sales@example.com", "url": "https://example.com/orders", "amount": 123.45,
              "percent": 12.5, "occurred": "2026-10-17T12:30:00"}
    saved = call(member, "POST", record_path, 201, {"values": values})
    detail = call(member, "GET", record_path + f"/{saved['id']}", 200)
    assert detail["values"]["description"] == values["description"]
    assert detail["values"]["amount"] == 123.45 and detail["values"]["occurred"] == values["occurred"]
    invalid = [
        ("phone", "abc"), ("email", "a@"), ("url", "javascript:alert(1)"),
        ("amount", 12.345), ("percent", 101), ("occurred", "2026-10-17T12:30:00.0001"),
        ("description", "x" * 10001),
    ]
    for key, value in invalid:
        rejected = {**values, key: value}
        call(member, "POST", record_path, 400, {"values": rejected})
    assert call(member, "GET", record_path, 200)["total"] == 1
    for key, value in (("phone", values["phone"]), ("amount", "123.45"),
                       ("percent", "12.5"), ("occurred", values["occurred"])):
        page = call(member, "GET", record_path + f"?filterKey={key}&filterValue={urllib.parse.quote(value)}", 200)
        assert page["total"] == 1, (key, page)
    call(member, "GET", record_path + "?filterKey=occurred&filterValue=bad", 400)
    call(member, "GET", record_path + "?filterKey=percent&filterValue=101", 400)
    call(member, "GET", record_path + "?filterKey=phone&filterValue=abc", 400)

    template = call(member, "GET", path + "/import-jobs/template", 200)
    book = openpyxl.load_workbook(io.BytesIO(template))
    keys = [book.active.cell(2, column).value for column in range(1, 10)]
    assert keys == ["record_no"] + [item[0] for item in definitions]
    row = [None, "Excel 订单", "导入行一\n导入行二", "13800138000", "buyer@example.com",
           "https://example.com/buyer", 88.25, "12%", datetime(2026, 10, 17, 15, 20, 30)]
    for column, value in enumerate(row, 1):
        book.active.cell(3, column, value)
    book.active.cell(3, 9).number_format = "yyyy-mm-dd hh:mm:ss"
    output = io.BytesIO()
    book.save(output)
    checked = preview(member, path + "/import-jobs", output.getvalue())
    assert checked["status"] == "READY", checked
    call(member, "POST", path + f"/import-jobs/{checked['id']}/commit", 200)
    page = call(member, "GET", record_path + "?filterKey=percent&filterValue=12", 200)
    assert page["total"] == 1 and page["items"][0]["values"]["title"] == "Excel 订单"
    assert page["items"][0]["values"]["occurred"] == "2026-10-17T15:20:30"

    export = call(member, "POST", path + "/export-jobs", 201,
                  {"columnKeys": [item[0] for item in definitions]})
    for _ in range(50):
        jobs = call(member, "GET", path + "/export-jobs", 200)
        job = next(item for item in jobs if item["id"] == export["id"])
        if job["status"] in ("SUCCEEDED", "FAILED"):
            break
        time.sleep(0.1)
    assert job["status"] == "SUCCEEDED" and job["rowCount"] == 2, job
    csv = call(member, "GET", path + f"/export-jobs/{job['id']}/download", 200).decode("utf-8-sig")
    assert "客户订单" in csv and "Excel 订单" in csv and "sales@example.com" in csv

    db = pymysql.connect(host="127.0.0.1", port=3306, user="root",
                         password=os.environ["APP_DB_PASSWORD"], database="unexamine_r15",
                         charset="utf8mb4", cursorclass=pymysql.cursors.DictCursor)
    try:
        with db.cursor() as cursor:
            cursor.execute("SELECT field_key,field_type FROM un_module_field WHERE module_id=%s ORDER BY sort_order", (mid,))
            assert [(item["field_key"], item["field_type"]) for item in cursor.fetchall()] == [
                (item[0], item[2]) for item in definitions]
            cursor.execute("SELECT COUNT(*) AS c FROM un_module_record WHERE module_id=%s", (mid,))
            assert cursor.fetchone()["c"] == 2
            cursor.execute("SELECT COUNT(*) AS c FROM un_module_record_index WHERE record_id=%s", (saved["id"],))
            assert cursor.fetchone()["c"] == 6
    finally:
        db.close()
    print(json.dumps({"status": "passed", "system_id": sid, "module_id": mid,
                      "checks": ["field_publish", "typed_create", "invalid_rejection", "filter",
                                 "excel_import", "percent_display", "datetime_import", "csv_export", "mysql_readback"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
